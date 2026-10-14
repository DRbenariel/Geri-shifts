"""Exports: Excel bytes + Google-Sheets tabs (appy.py:587-945, 972-1356).

Excel builders return (bytes, filename) — the API streams them as downloads.
"""
from __future__ import annotations

import calendar
import io
from datetime import date, timedelta

import pandas as pd

from .constants import (BATCHED_ABSENT_STATUSES, DAILY_DEPTS_ALL, GRID_STATUS_LABEL_SHORT, PNIM_DEPT,
                        PNIM_ICONS, PNIM_SIDES, XL_STATUS_FILL)
from .daily import (derive_auto_status, get_fri_shift_workers, get_night_duty, incoming_transfers,
                    night_shift_dept_label, personal_status, wsd_get_note, wsd_get_status)
from .days import is_functional_friday, is_functional_saturday
from .people import get_dept_managers, sort_employees_by_role

WD = ["א", "ב", "ג", "ד", "ה", "ו", "ש"]


def _year_of(year_month, year=None):
    if year is not None:
        return year
    try:
        return int(str(year_month).split('-')[0])
    except Exception:
        return 2026


def _safe(dept_name, n):
    return dept_name.replace("'", "").replace('"', '').replace('/', '_')[:n]


def wsd_sheet_name(dept_name, year_month):
    return f"WSD_{_safe(dept_name, 25)}_{year_month}"


# ── night schedule → Schedule_Export ─────────────────────────────────────────
def schedule_wide_df(ctx, month, year=2026) -> pd.DataFrame:
    rows = []
    sched = ctx.schedule
    for d in range(1, calendar.monthrange(year, month)[1] + 1):
        d_obj = date(year, month, d)
        d_str = str(d_obj)
        row = {'תאריך': d_str, 'יום': ["ב'", "ג'", "ד'", "ה'", "ו'", "ש'", "א'"][d_obj.weekday()],
               'פנימית גריאטרית': '', 'שיקום': '',
               'שישי בוקר - פנימית (1)': '', 'שישי בוקר - פנימית (2)': '',
               'שישי בוקר - שיקום (1)': '', 'שישי בוקר - שיקום (2)': ''}
        for _, s in sched[sched['date'] == d_str].iterrows():
            if s['employee'] == '---':
                continue
            if s['dept'] in row:
                row[s['dept']] = s['employee']
        rows.append(row)
    return pd.DataFrame(rows)[['תאריך', 'יום', 'פנימית גריאטרית', 'שיקום',
                               'שישי בוקר - פנימית (1)', 'שישי בוקר - פנימית (2)',
                               'שישי בוקר - שיקום (1)', 'שישי בוקר - שיקום (2)']]


def export_schedule_wide(ctx, month, year=2026):
    ctx.save("Schedule_Export", schedule_wide_df(ctx, month, year), is_rtl=True)


# ── batched per-day data (shared by Excel + Sheets exports) ──────────────────
def build_batched_day_data(ctx, dept_name, year_month, view_month, year=None):
    year = _year_of(year_month, year)
    days = list(range(1, calendar.monthrange(year, view_month)[1] + 1))
    is_pnim = dept_name == PNIM_DEPT
    dr = ctx.dept_rotation
    if dr.empty or 'employee' not in dr.columns:
        return None, None, None, None, None
    mask = (dr['year_month'].astype(str) == year_month) & (dr['daily_dept'].astype(str) == dept_name)
    employees = dr[mask]['employee'].astype(str).str.strip().tolist()
    employees = employees + [m for m in get_dept_managers(ctx, dept_name) if m not in employees]
    side_map = {}
    if is_pnim and 'side' in dr.columns:
        side_map = (dr[dr['year_month'].astype(str) == year_month].set_index('employee')['side']
                    .fillna('').astype(str).str.strip().to_dict())
    incoming = incoming_transfers(ctx, dept_name, year_month)
    transfer_side = {e: r.get('side', '') for e, r in incoming.items()}
    if not employees and not incoming:
        return None, None, None, None, None
    employees = sort_employees_by_role(ctx, employees)

    def emp_side(e):
        return side_map.get(e, '') or transfer_side.get(e, '')

    pw, pt, pa = {}, {}, {}
    pws = {} if is_pnim else None
    for d in days:
        ds = f"{year}-{view_month:02d}-{d:02d}"
        do = date(year, view_month, d)
        is_sat = is_functional_saturday(ctx, do)
        is_fri = (not is_sat) and is_functional_friday(ctx, do)
        workers, absent = [], []
        if is_sat:
            pass
        elif is_fri:
            workers = list(get_fri_shift_workers(ctx, ds, dept_name))
        else:
            for emp in employees:
                status = derive_auto_status(ctx, ds, emp, daily_dept=dept_name)
                if status == "":
                    continue
                if status in BATCHED_ABSENT_STATUSES:
                    if status == "אחרי תורנות":
                        lbl = night_shift_dept_label(ctx, (do - timedelta(days=1)).strftime("%Y-%m-%d"), emp)
                        absent.append(f"{emp} - אחרי תורנות ({lbl})" if lbl else f"{emp} - אחרי תורנות")
                    else:
                        absent.append(f"{emp} - {status}")
                else:
                    workers.append(emp)
        for emp, rec in incoming.items():
            if d in rec['days'] and emp not in workers:
                status = derive_auto_status(ctx, ds, emp, daily_dept=dept_name)
                if status and status not in BATCHED_ABSENT_STATUSES:
                    workers.append(emp)
        workers = sort_employees_by_role(ctx, workers)
        pw[d] = workers
        pt[d] = get_night_duty(ctx, ds, dept_name)
        pa[d] = absent
        if is_pnim:
            by_side = {s: [] for s in PNIM_SIDES}
            for emp in workers:
                s = emp_side(emp)
                if s in by_side:
                    by_side[s].append(emp)
            pws[d] = by_side
    return employees, pw, pt, pa, pws


def dept_sheet_df(ctx, dept_name, year_month, view_month, year=None):
    """Wide DataFrame for the WSD_<dept>_<ym> Google-Sheets tab (appy._export_dept_grid)."""
    year = _year_of(year_month, year)
    days = list(range(1, calendar.monthrange(year, view_month)[1] + 1))
    res = build_batched_day_data(ctx, dept_name, year_month, view_month, year=year)
    if not res or res[0] is None:
        return None
    _e, pw, pt, pa, pws = res
    toranet = "תורן פנגר" if dept_name == "פנימית גריאטרית" else "תורן שיקום"
    max_w = max((len(v) for v in pw.values()), default=1)
    max_a = max((len(v) for v in pa.values()), default=0)
    cols = ['#'] + [str(d) for d in days]
    rows = [{'#': f"{dept_name} — {year_month}",
             **{str(d): f"{d} / {WD[(date(year, view_month, d).weekday() + 1) % 7]}" for d in days}}]

    def block(getter, n):
        for i in range(n):
            r = {'#': ''}
            for d in days:
                lst = getter(d)
                r[str(d)] = lst[i] if i < len(lst) else ''
            rows.append(r)

    if pws is not None:
        for s in PNIM_SIDES:
            mx = max((len(v.get(s, [])) for v in pws.values()), default=0)
            rows.append({'#': f"{PNIM_ICONS[s]} צד {s}", **{str(d): '' for d in days}})
            block(lambda d, s=s: pws.get(d, {}).get(s, []), max(mx, 1))
    else:
        rows.append({'#': '— עובדים —', **{str(d): '' for d in days}})
        block(lambda d: pw.get(d, []), max_w)
    rows.append({'#': toranet, **{str(d): pt.get(d) or '' for d in days}})
    if max_a > 0:
        rows.append({'#': '— לא נמצאים —', **{str(d): '' for d in days}})
        block(lambda d: pa.get(d, []), max_a)
    return pd.DataFrame(rows, columns=cols)


def export_dept_to_sheet(ctx, dept_name, year_month, view_month, year=None, share_anyone=False):
    """Write WSD_<dept>_<ym>, set RTL, return the tab URL. `share_anyone` reproduces the
    Streamlit 'anyone with the link can view' behaviour — off by default (decision pending)."""
    df = dept_sheet_df(ctx, dept_name, year_month, view_month, year)
    if df is None:
        raise ValueError("שגיאה בייצוא — וודא שיש שיבוצים למחלקה זו")
    name = wsd_sheet_name(dept_name, year_month)
    ctx.save(name, df)
    try:
        ctx.store.set_rtl(name)
    except Exception:
        pass
    if share_anyone and hasattr(ctx.store, '_spreadsheet'):
        try:
            ctx.store._spreadsheet().share('', perm_type='anyone', role='reader')
        except Exception:
            pass
    return ctx.store.tab_url(name)


# ── Excel ────────────────────────────────────────────────────────────────────
def write_batched_sheet(ws, dept_name, year_month, view_month, pw, pt, pa, pws=None, year=None):
    """Week-block, print-ready layout (appy._write_batched_sheet) — verbatim port."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.properties import PageSetupProperties

    year = _year_of(year_month, year)
    toranet = "תורן פנגר" if dept_name == "פנימית גריאטרית" else "תורן שיקום"
    thin = Side(style='thin', color='CBD5E1')
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_hdr_dept = PatternFill("solid", fgColor="1E293B")
    f_hdr_day = PatternFill("solid", fgColor="334155")
    f_sec_work = PatternFill("solid", fgColor="DCFCE7")
    f_sec_toran = PatternFill("solid", fgColor="EDE9FE")
    f_sec_absent = PatternFill("solid", fgColor="FEF2F2")
    f_work_cell = PatternFill("solid", fgColor="F0FDF4")
    f_fri_cell = PatternFill("solid", fgColor="FFFBEB")
    f_sat_cell = PatternFill("solid", fgColor="F1F5F9")
    f_abs_cell = PatternFill("solid", fgColor="FFF5F5")
    ws.sheet_view.rightToLeft = True

    def _set(row, col, val="", **kw):
        c = ws.cell(row, col, val)
        c.border = bdr
        for a, v in kw.items():
            setattr(c, a, v)
        return c

    TOTAL = 8
    _set(1, 1, f"{dept_name} — {year_month}", font=Font(bold=True, size=12, color="FFFFFF"),
         fill=f_hdr_dept, alignment=Alignment(horizontal='right'))
    for c in range(2, TOTAL + 1):
        _set(1, c, fill=f_hdr_dept)
    weeks = calendar.Calendar(firstweekday=6).monthdayscalendar(year, view_month)

    def wmax(wk):
        return max((len(pw.get(d, [])) for d in wk if d), default=1)

    def smax(wk, s):
        return max((len(pws.get(d, {}).get(s, [])) for d in wk if d), default=0)

    def amax(wk):
        return max((len(pa.get(d, [])) for d in wk if d), default=0)

    def block_rows(wk):
        n = 1
        if pws is not None:
            for s in PNIM_SIDES:
                n += 1 + max(smax(wk, s), 1)
        else:
            n += 1 + max(wmax(wk), 1)
        return n + 1 + 1 + amax(wk)

    def day_fill(j):
        return f_sat_cell if j == 6 else f_fri_cell if j == 5 else f_work_cell

    r = 2

    def write_week(wk):
        nonlocal r
        real = [d for d in wk if d]
        rng = f"{real[0]}–{real[-1]}/{view_month}" if real else ""
        _set(r, 1, rng, font=Font(bold=True, color="FFFFFF", size=9), fill=f_hdr_day,
             alignment=Alignment(horizontal='center', vertical='center'))
        for j in range(7):
            d = wk[j]
            _set(r, 2 + j, f"{d}\n{WD[j]}" if d else "", font=Font(bold=True, color="FFFFFF", size=9),
                 fill=f_hdr_day, alignment=Alignment(horizontal='center', vertical='center', wrap_text=True))
        ws.row_dimensions[r].height = 26
        r += 1

        def worker_rows(getter, n):
            nonlocal r
            for i in range(n):
                _set(r, 1)
                for j in range(7):
                    d = wk[j]
                    lst = getter(d) if d else []
                    _set(r, 2 + j, lst[i] if i < len(lst) else "", font=Font(size=9), fill=day_fill(j),
                         alignment=Alignment(horizontal='right', vertical='center'))
                r += 1

        if pws is not None:
            f_side = {"ורוד": PatternFill("solid", fgColor="FCE7F3"), "כחול": PatternFill("solid", fgColor="DBEAFE")}
            for s in PNIM_SIDES:
                _set(r, 1, f"{PNIM_ICONS[s]} צד {s}", font=Font(bold=True, size=9, color="1E293B"),
                     fill=f_side.get(s, f_sec_work), alignment=Alignment(horizontal='right'))
                for j in range(7):
                    _set(r, 2 + j, fill=f_side.get(s, f_sec_work))
                r += 1
                worker_rows(lambda d, s=s: pws.get(d, {}).get(s, []), max(smax(wk, s), 1))
        else:
            _set(r, 1, "עובדים", font=Font(bold=True, size=9, color="166534"), fill=f_sec_work,
                 alignment=Alignment(horizontal='right'))
            for j in range(7):
                _set(r, 2 + j, fill=f_sec_work)
            r += 1
            worker_rows(lambda d: pw.get(d, []), max(wmax(wk), 1))
        _set(r, 1, toranet, font=Font(bold=True, size=9, color="5B21B6"), fill=f_sec_toran,
             alignment=Alignment(horizontal='right'))
        for j in range(7):
            d = wk[j]
            _set(r, 2 + j, (pt.get(d) or "") if d else "", font=Font(size=9, color="5B21B6"), fill=f_sec_toran,
                 alignment=Alignment(horizontal='right', vertical='center'))
        r += 1
        _set(r, 1, "לא נמצאים", font=Font(bold=True, size=9, color="991B1B"), fill=f_sec_absent,
             alignment=Alignment(horizontal='right'))
        for j in range(7):
            _set(r, 2 + j, fill=f_sec_absent)
        r += 1
        for i in range(amax(wk)):
            _set(r, 1)
            for j in range(7):
                d = wk[j]
                lst = pa.get(d, []) if d else []
                val = lst[i] if i < len(lst) else ""
                _set(r, 2 + j, val, font=Font(size=9, color="991B1B" if val else "9CA3AF"), fill=f_abs_cell,
                     alignment=Alignment(horizontal='right', vertical='center'))
            r += 1

    BUDGET = 34
    page = 0
    for wk in weeks:
        b = block_rows(wk)
        if page > 0 and page + b + 1 > BUDGET:
            ws.row_breaks.append(Break(id=r - 1))
            page = 0
        write_week(wk)
        r += 1
        page += b + 1
    ws.column_dimensions['A'].width = 15
    for j in range(7):
        ws.column_dimensions[get_column_letter(2 + j)].width = 13
    ws.freeze_panes = "B2"
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.2, footer=0.2)
    ws.print_title_rows = '1:1'


def _wb_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def dept_batched_excel(ctx, dept_name, year_month, view_month, year=None):
    import openpyxl
    try:
        e, pw, pt, pa, pws = build_batched_day_data(ctx, dept_name, year_month, view_month, year=year)
        if e is None:
            return None, None
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = _safe(dept_name, 31)
        write_batched_sheet(ws, dept_name, year_month, view_month, pw, pt, pa, pws, year=year)
        return _wb_bytes(wb), f"סידור_{_safe(dept_name, 25)}_{year_month}.xlsx"
    except Exception:
        return None, None


def all_depts_batched_excel(ctx, year_month, view_month, year=None):
    import openpyxl
    try:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        added = False
        for dept in DAILY_DEPTS_ALL:
            e, pw, pt, pa, pws = build_batched_day_data(ctx, dept, year_month, view_month, year=year)
            if e is None:
                continue
            ws = wb.create_sheet(title=_safe(dept, 31))
            write_batched_sheet(ws, dept, year_month, view_month, pw, pt, pa, pws, year=year)
            added = True
        if not added:
            return None, None
        return _wb_bytes(wb), f"סידור_כל_המחלקות_{year_month}.xlsx"
    except Exception:
        return None, None


def dept_grid_excel(ctx, dept_name, year_month, view_month, year=2026):
    """Colour-coded employee × day grid (appy._export_dept_grid_excel)."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    try:
        days = list(range(1, calendar.monthrange(year, view_month)[1] + 1))
        dr = ctx.dept_rotation
        if dr.empty or 'employee' not in dr.columns:
            return None, None
        mask = (dr['year_month'].astype(str) == year_month) & (dr['daily_dept'].astype(str) == dept_name)
        employees = dr[mask]['employee'].astype(str).str.strip().tolist()
        if not employees:
            return None, None
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = dept_name[:31]
        ws.sheet_view.rightToLeft = True
        thin = Side(style='thin', color='CBD5E1')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.cell(1, 1, f"{dept_name} — {year_month}").font = Font(bold=True, size=12)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(days) + 1)
        ws.cell(2, 1, "עובד/ת").font = Font(bold=True)
        for i, d in enumerate(days):
            c = ws.cell(2, i + 2, f"{d}\n{WD[(date(year, view_month, d).weekday() + 1) % 7]}")
            c.font = Font(bold=True, size=9)
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = border
        ws.row_dimensions[2].height = 28
        ws.column_dimensions['A'].width = 16
        for ri, emp in enumerate(employees):
            rn = ri + 3
            ws.cell(rn, 1, emp).font = Font(bold=False, size=10)
            ws.cell(rn, 1).alignment = Alignment(horizontal='right')
            for i, d in enumerate(days):
                ds = f"{year}-{view_month:02d}-{d:02d}"
                raw = wsd_get_status(ctx, ds, emp, default=None)
                status = raw if raw is not None else derive_auto_status(ctx, ds, emp, daily_dept=dept_name)
                c = ws.cell(rn, i + 2)
                c.value = GRID_STATUS_LABEL_SHORT.get(status, status)
                c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True, shrink_to_fit=False)
                c.fill = PatternFill("solid", fgColor=XL_STATUS_FILL.get(status, "F8FAFC"))
                c.border = border
                c.font = Font(size=9)
                ws.column_dimensions[openpyxl.utils.get_column_letter(i + 2)].width = 6
        ws.freeze_panes = "B3"
        return _wb_bytes(wb), f"WSD_{_safe(dept_name, 25)}_{year_month}.xlsx"
    except Exception:
        return None, None


def personal_excel(ctx, user_name, view_month, daily_dept=None, year=2026):
    """One employee's month, one row per day (appy._export_personal_schedule_excel)."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    try:
        WD_FULL = ["ראשון", "שני", "שלישי", "רביעי", "חמישי", "שישי", "שבת"]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "לוח עבודה אישי"
        ws.sheet_view.rightToLeft = True
        thin = Side(style='thin', color='CBD5E1')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for ci, h in enumerate(["תאריך", "יום", "סטטוס", "הערה"]):
            c = ws.cell(1, ci + 1, h)
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor="334155")
            c.alignment = Alignment(horizontal='center')
            c.border = border
        for col, w in zip("ABCD", (12, 10, 14, 22)):
            ws.column_dimensions[col].width = w
        for d in range(1, calendar.monthrange(year, view_month)[1] + 1):
            ds = f"{year}-{view_month:02d}-{d:02d}"
            status = personal_status(ctx, ds, user_name, daily_dept=daily_dept)
            rn = d + 1
            ws.cell(rn, 1, ds).border = border
            ws.cell(rn, 2, WD_FULL[(date(year, view_month, d).weekday() + 1) % 7]).border = border
            sc = ws.cell(rn, 3, status)
            sc.fill = PatternFill("solid", fgColor=XL_STATUS_FILL.get(status, "F8FAFC"))
            sc.alignment = Alignment(horizontal='center')
            sc.border = border
            ws.cell(rn, 4, wsd_get_note(ctx, ds, user_name)).border = border
        ws.freeze_panes = "A2"
        safe_name = str(user_name).replace(' ', '_')[:20]
        return _wb_bytes(wb), f"סידור_{safe_name}_{year}-{view_month:02d}.xlsx"
    except Exception:
        return None, None
